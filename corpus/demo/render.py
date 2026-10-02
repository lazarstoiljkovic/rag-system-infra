"""
Renderovanje sintetickih dokumenata u PDF i DOCX, iz istog opisa.

Dokument je lista blokova:

  ("h1", tekst)        naslov dokumenta
  ("h2", tekst)        naslov odeljka
  ("p", tekst)         pasus
  ("table", {...})     tabela: caption, rows (prvi red je zaglavlje), widths
  ("figure", {...})    slika: caption i spec dijagrama ili grafikona

Isti opis daje PDF (PyMuPDF) ili DOCX (python-docx), pa obe grane ekstraktora
dobijaju uporediv sadrzaj.

Pravila koja proizilaze iz ekstraktora (`lambda/extract_and_prepare/extractor.py`):
  - tabele u PDF-u imaju punu mrezu linija, jer ih `find_tables()` trazi po
    linijama;
  - slike su RASTERSKE (PNG), jer `get_images()` ne vidi vektorske crteze.
"""

import math
import os

import fitz
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle  # noqa: E402

_TTF = os.path.join(matplotlib.get_data_path(), "fonts", "ttf")
FONT = os.path.join(_TTF, "DejaVuSans.ttf")
FONT_BOLD = os.path.join(_TTF, "DejaVuSans-Bold.ttf")

_INK = "#1f3b63"
_FILL = "#e8eef7"
_BAR = "#3b6ea5"


# --- slike ----------------------------------------------------------------

def _border_point(center, toward, half_w, half_h, pad=0.8):
    """Tacka na ivici pravougaonika oko `center`, u pravcu tacke `toward`."""
    dx, dy = toward[0] - center[0], toward[1] - center[1]
    if dx == 0 and dy == 0:
        return center
    t = min(half_w / abs(dx) if dx else math.inf, half_h / abs(dy) if dy else math.inf)
    length = math.hypot(dx, dy)
    return (center[0] + dx * t + dx / length * pad, center[1] + dy * t + dy / length * pad)


def _beside(start, end, box_h, distance=3.4):
    """
    Tacka pored sredine linije: iznad nje, a kod uspravne linije desno.

    Kod VODORAVNE linije natpis ide iznad gornje ivice kutija: izmedju dve
    kutije u redu obicno nema mesta za natpis, pa bi presao preko njihovih
    ivica i delimicno sakrio tekst u kutiji.
    """
    dx, dy = end[0] - start[0], end[1] - start[1]
    length = math.hypot(dx, dy) or 1.0
    if abs(dy) < abs(dx) * 0.2:
        return ((start[0] + end[0]) / 2, (start[1] + end[1]) / 2 + box_h / 2 + 2.3)
    nx, ny = -dy / length, dx / length
    if ny < 0 or (abs(ny) < 1e-9 and nx < 0):
        nx, ny = -nx, -ny
    if abs(dx) < abs(dy) * 0.2:          # skoro uspravna linija: natpis desno
        nx, ny = abs(nx) or 1.0, 0.0
        distance *= 2.2
    return ((start[0] + end[0]) / 2 + nx * distance,
            (start[1] + end[1]) / 2 + ny * distance)


def draw_diagram(spec, path):
    """
    Dijagram arhitekture. spec:
      size   (sirina, visina) u jedinicama crteza
      box    (sirina, visina) kutije
      nodes  {id: (x_centra, y_centra, natpis)}
      edges  [(od, ka, natpis) ili (od, ka, natpis, (x_natpisa, y_natpisa))]
      groups [(x0, y0, x1, y1, natpis) ili (..., natpis, "dole")] — isprekidane
             granice (zona, klaster); "dole" stavlja natpis uz donju ivicu

    Natpis strelice podrazumevano stoji PORED linije, ne na njoj: natpis sa
    belom pozadinom preko kratke strelice sakrije i nju i njen vrh, pa se smer
    ne vidi — a smer je upravo ono sto caption mora da procita.
    """
    width, height = spec["size"]
    box_w, box_h = spec.get("box", (24, 9))
    fig, ax = plt.subplots(figsize=(11, 11 * height / width), dpi=150)
    ax.set_xlim(0, width)
    ax.set_ylim(0, height)
    ax.axis("off")

    for group in spec.get("groups", []):
        x0, y0, x1, y1, label = group[:5]
        ax.add_patch(Rectangle((x0, y0), x1 - x0, y1 - y0, fill=False,
                               linestyle="--", linewidth=1.2, edgecolor="#555555"))
        bottom = len(group) > 5 and group[5] == "dole"
        ax.text(x0 + 1.5, y0 + 1.5 if bottom else y1 - 3, label, fontsize=9.5,
                style="italic", va="bottom" if bottom else "baseline")

    for x, y, label in spec["nodes"].values():
        ax.add_patch(FancyBboxPatch((x - box_w / 2, y - box_h / 2), box_w, box_h,
                                    boxstyle="round,pad=0.4", facecolor=_FILL,
                                    edgecolor=_INK, linewidth=1.5))
        ax.text(x, y, label, ha="center", va="center", fontsize=9.5)

    for edge in spec["edges"]:
        src, dst, label = edge[:3]
        a, b = spec["nodes"][src][:2], spec["nodes"][dst][:2]
        start = _border_point(a, b, box_w / 2, box_h / 2)
        end = _border_point(b, a, box_w / 2, box_h / 2)
        ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=15,
                                     linewidth=1.5, color=_INK))
        if label:
            lx, ly = edge[3] if len(edge) > 3 else _beside(start, end, box_h)
            ax.text(lx, ly, label, fontsize=8.3, ha="center", va="center",
                    bbox=dict(facecolor="white", edgecolor="none", pad=1))

    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def draw_chart(spec, path):
    """
    Grafikon sa TACNIM vrednostima ispisanim uz svaku tacku ili stubic. spec:
      kind        "bar" ili "line"
      title, x_label, y_label, categories, values, y_max, fmt
      y_min       opciono; za vrednosti blizu jedna drugoj (npr. dostupnost
                  99,8%), inace bi svi stubici izgledali isto
    """
    fig, ax = plt.subplots(figsize=(7, 4), dpi=150)
    cats, vals = spec["categories"], spec["values"]
    fmt = spec.get("fmt", "{:.1f}")
    y_min = spec.get("y_min", 0)
    offset = (spec["y_max"] - y_min) * 0.015

    if spec["kind"] == "bar":
        bars = ax.bar(cats, [v - y_min for v in vals], bottom=y_min, color=_BAR)
        for bar, value in zip(bars, vals):
            ax.text(bar.get_x() + bar.get_width() / 2, value + offset, fmt.format(value),
                    ha="center", fontsize=9.5)
    else:
        ax.plot(cats, vals, marker="o", color=_BAR, linewidth=2)
        for x, value in zip(cats, vals):
            ax.annotate(fmt.format(value), (x, value), textcoords="offset points",
                        xytext=(0, 7), ha="center", fontsize=9.5)

    ax.set_title(spec["title"])
    ax.set_xlabel(spec["x_label"])
    ax.set_ylabel(spec["y_label"])
    ax.set_ylim(y_min, spec["y_max"])
    ax.spines[["top", "right"]].set_visible(False)
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def draw_sequence(spec, path):
    """
    Dijagram sekvence. spec:
      participants  [natpis, ...] — redom, sleva nadesno
      messages      [(od, ka, natpis) ili (od, ka, natpis, "povratna")]
                    od/ka su indeksi u `participants`; "povratna" crta
                    isprekidanu strelicu (odgovor)

    Poruke se numerisu redom ("1. ..."), jer je redosled upravo ono sto
    pitanja traze, a bez brojeva bi ga caption morao da zakljuci iz polozaja.
    """
    parts = spec["participants"]
    msgs = spec["messages"]
    step = 30
    width = step * len(parts)
    height = 18 + 9 * len(msgs)
    fig, ax = plt.subplots(figsize=(11, 11 * height / width), dpi=150)
    ax.set_xlim(0, width)
    ax.set_ylim(0, height)
    ax.axis("off")
    xs = [step / 2 + i * step for i in range(len(parts))]
    top = height - 6
    for x, label in zip(xs, parts):
        ax.add_patch(FancyBboxPatch((x - 12, top - 4), 24, 8, boxstyle="round,pad=0.3",
                                    facecolor=_FILL, edgecolor=_INK, linewidth=1.5))
        ax.text(x, top, label, ha="center", va="center", fontsize=9)
        ax.plot([x, x], [top - 4.5, 2], linestyle=(0, (3, 3)), color="#888888", linewidth=1)
    for n, msg in enumerate(msgs, start=1):
        src, dst, label = msg[:3]
        back = len(msg) > 3 and msg[3] == "povratna"
        y = top - 9 - (n - 1) * 9
        ax.add_patch(FancyArrowPatch((xs[src], y), (xs[dst], y), arrowstyle="-|>",
                                     mutation_scale=13, linewidth=1.4, color=_INK,
                                     linestyle="--" if back else "-"))
        text = "{}. {}".format(n, label) if label else "{}.".format(n)
        ax.text((xs[src] + xs[dst]) / 2, y + 1.6, text, ha="center", va="bottom",
                fontsize=8.3, bbox=dict(facecolor="white", edgecolor="none", pad=0.5))
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


_SERIES = ["#3b6ea5", "#d98b2b", "#4f9d69", "#b04a5a", "#7a5aa6"]


def draw_multi_chart(spec, path):
    """
    Grafikoni sa vise serija ili delova. spec["kind"]:
      multiline  categories, series {naziv: vrednosti}; vrednosti NISU ispisane
                 (citaju se sa ose) — namerno teze za opis recima
      grouped    categories, series {naziv: vrednosti}; vrednosti ispisane
      pie        labels, values (procenti, zbir 100)
      gantt      months [natpisi], tasks [(naziv, prvi_mesec, poslednji_mesec)]
                 indeksi meseca od 0; trajanje se cita sa ose, bez brojeva
    """
    kind = spec["kind"]
    fmt = spec.get("fmt", "{:.1f}")
    if kind == "pie":
        fig, ax = plt.subplots(figsize=(6.5, 4.5), dpi=150)
        ax.pie(spec["values"], labels=spec["labels"], colors=_SERIES + ["#999999"],
               autopct=lambda p: fmt.format(p) + "%", startangle=90, counterclock=False,
               textprops={"fontsize": 9.5})
        ax.set_title(spec["title"])
        ax.axis("equal")
    elif kind == "gantt":
        tasks, months = spec["tasks"], spec["months"]
        fig, ax = plt.subplots(figsize=(8, 0.5 * len(tasks) + 1.6), dpi=150)
        for row, (name, first, last) in enumerate(tasks):
            ax.barh(row, last - first + 1, left=first - 0.5, height=0.55, color=_BAR)
        ax.set_yticks(range(len(tasks)))
        ax.set_yticklabels([t[0] for t in tasks], fontsize=9.5)
        ax.invert_yaxis()
        ax.set_xticks(range(len(months)))
        ax.set_xticklabels(months, fontsize=9)
        ax.set_xlim(-0.5, len(months) - 0.5)
        ax.grid(axis="x", linestyle=":", color="#bbbbbb")
        ax.set_axisbelow(True)
        ax.set_title(spec["title"])
        ax.set_xlabel(spec.get("x_label", ""))
        ax.spines[["top", "right"]].set_visible(False)
    else:
        cats, series = spec["categories"], spec["series"]
        fig, ax = plt.subplots(figsize=(7.5, 4.2), dpi=150)
        if kind == "multiline":
            for color, (name, vals) in zip(_SERIES, series.items()):
                ax.plot(cats, vals, marker="o", linewidth=2, color=color, label=name)
            ax.grid(axis="y", linestyle=":", color="#bbbbbb")
        elif kind == "grouped":
            n = len(series)
            width = 0.8 / n
            for i, (color, (name, vals)) in enumerate(zip(_SERIES, series.items())):
                xs = [c + (i - (n - 1) / 2) * width for c in range(len(cats))]
                bars = ax.bar(xs, vals, width=width, color=color, label=name)
                for bar, value in zip(bars, vals):
                    ax.text(bar.get_x() + bar.get_width() / 2, value + spec["y_max"] * 0.012,
                            fmt.format(value), ha="center", fontsize=8.5)
            ax.set_xticks(range(len(cats)))
            ax.set_xticklabels(cats)
        else:
            raise ValueError("nepoznata vrsta grafikona: {}".format(kind))
        ax.set_ylim(spec.get("y_min", 0), spec["y_max"])
        ax.set_title(spec["title"])
        ax.set_xlabel(spec["x_label"])
        ax.set_ylabel(spec["y_label"])
        ax.legend(fontsize=8.5, frameon=False)
        ax.spines[["top", "right"]].set_visible(False)
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)


def render_figure(figure, path):
    spec = figure["spec"]
    if "nodes" in spec:
        draw_diagram(spec, path)
    elif "participants" in spec:
        draw_sequence(spec, path)
    elif spec.get("kind") in ("multiline", "grouped", "pie", "gantt"):
        draw_multi_chart(spec, path)
    else:
        draw_chart(spec, path)


# --- PDF ------------------------------------------------------------------

_MARGIN = 56
_PAGE_W, _PAGE_H = fitz.paper_size("a4")
_STYLES = {"h1": (17, True, 10), "h2": (12.5, True, 6), "p": (10.5, False, 8)}


class _PdfFlow:
    """Tekuca pozicija na strani; nova strana kad blok ne staje."""

    def __init__(self):
        self.doc = fitz.open()
        self.font = fitz.Font(fontfile=FONT)
        self.font_bold = fitz.Font(fontfile=FONT_BOLD)
        self._new_page()

    def _new_page(self):
        self.page = self.doc.new_page(width=_PAGE_W, height=_PAGE_H)
        self.y = _MARGIN

    def _bottom(self):
        return _PAGE_H - _MARGIN

    def text(self, text, size, bold, gap_after, x0=_MARGIN, x1=_PAGE_W - _MARGIN):
        """Pasus sa prelamanjem; ako ne staje, ide na novu stranu (ne deli se)."""
        for attempt in range(2):
            rect = fitz.Rect(x0, self.y, x1, self._bottom())
            rest = self.page.insert_textbox(
                rect, text, fontsize=size, fontname="dvb" if bold else "dv",
                fontfile=FONT_BOLD if bold else FONT)
            if rest >= 0:
                self.y = rect.y1 - rest + gap_after
                return
            if attempt == 0:
                self._new_page()
        raise ValueError("blok ne staje ni na praznu stranu: {}...".format(text[:50]))

    def _lines(self, text, width, size, bold):
        font = self.font_bold if bold else self.font
        return sum(max(1, math.ceil(font.text_length(part, fontsize=size) / width))
                   for part in text.split("\n"))

    def table(self, table):
        size = 9.3
        widths = table["widths"]
        heights = [
            max(self._lines(str(cell), widths[c] - 8, size, r == 0)
                for c, cell in enumerate(row)) * size * 1.3 + 8
            for r, row in enumerate(table["rows"])
        ]
        if self.y + 22 + sum(heights) > self._bottom():
            self._new_page()
        self.text(table["caption"], 10, True, 4)

        x0, y = _MARGIN, self.y
        x1 = x0 + sum(widths)
        y_end = y + sum(heights)
        self.page.draw_line((x0, y), (x1, y), width=0.8)
        for r, row in enumerate(table["rows"]):
            x = x0
            for c, cell in enumerate(row):
                rect = fitz.Rect(x + 4, y + 4, x + widths[c] - 4, y + heights[r])
                rest = self.page.insert_textbox(
                    rect, str(cell), fontsize=size, fontname="dvb" if r == 0 else "dv",
                    fontfile=FONT_BOLD if r == 0 else FONT)
                if rest < 0:
                    raise ValueError("celija ne staje: {}".format(cell))
                x += widths[c]
            y += heights[r]
            self.page.draw_line((x0, y), (x1, y), width=0.8)
        x = x0
        for w in list(widths) + [0]:
            self.page.draw_line((x, self.y), (x, y_end), width=0.8)
            x += w
        self.y = y_end + 14

    def figure(self, figure, png):
        pix = fitz.Pixmap(png)
        width = (_PAGE_W - 2 * _MARGIN) * figure.get("scale", 1.0)
        height = width * pix.height / pix.width
        if self.y + height + 24 > self._bottom():
            self._new_page()
        x0 = (_PAGE_W - width) / 2
        self.page.insert_image(fitz.Rect(x0, self.y, x0 + width, self.y + height), filename=png)
        self.y += height + 4
        self.text(figure["caption"], 9.5, True, 12)


def render_pdf(blocks, path, workdir, title, author):
    flow = _PdfFlow()
    figure_no = 0
    for kind, content in blocks:
        if kind in _STYLES:
            size, bold, gap = _STYLES[kind]
            flow.text(content, size, bold, gap)
        elif kind == "table":
            flow.table(content)
        elif kind == "figure":
            figure_no += 1
            png = os.path.join(workdir, "{}-{}.png".format(os.path.basename(path), figure_no))
            render_figure(content, png)
            flow.figure(content, png)
        else:
            raise ValueError("nepoznat blok: {}".format(kind))
    flow.doc.set_metadata({"title": title, "author": author})
    flow.doc.save(path, garbage=3, deflate=True)
    flow.doc.close()


# --- DOCX -----------------------------------------------------------------

def render_docx(blocks, path, workdir, title, author):
    import docx
    from docx.shared import Cm, Pt

    document = docx.Document()
    document.core_properties.title = title
    document.core_properties.author = author
    figure_no = 0
    for kind, content in blocks:
        if kind == "h1":
            document.add_heading(content, level=0)
        elif kind == "h2":
            document.add_heading(content, level=1)
        elif kind == "p":
            document.add_paragraph(content)
        elif kind == "table":
            caption = document.add_paragraph()
            caption.add_run(content["caption"]).bold = True
            rows = content["rows"]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    paragraph = table.cell(r, c).paragraphs[0]
                    run = paragraph.add_run(str(cell))
                    run.bold = r == 0
                    run.font.size = Pt(9.5)
        elif kind == "figure":
            figure_no += 1
            png = os.path.join(workdir, "{}-{}.png".format(os.path.basename(path), figure_no))
            render_figure(content, png)
            document.add_picture(png, width=Cm(16 * content.get("scale", 1.0)))
            caption = document.add_paragraph()
            caption.add_run(content["caption"]).bold = True
        else:
            raise ValueError("nepoznat blok: {}".format(kind))
    document.save(path)


# --- citanje nazad, za proveru ---------------------------------------------

def extracted_text(path):
    """
    Sav TEKST dokumenta onako kako ga vidi ekstrakcija — bez slika.

    Sluzi proveri pitanja: odgovor oznacen kao "samo na slici" ne sme se naci
    nigde ovde, inace bi i varijanta A mogla da odgovori, pa demo razlike
    izmedju A i B ne bi vazio.
    """
    if path.endswith(".pdf"):
        with fitz.open(path) as doc:
            return "\n".join(page.get_text() for page in doc)
    import docx
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
