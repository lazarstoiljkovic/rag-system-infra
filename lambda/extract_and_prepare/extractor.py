"""
Ekstrakcija elemenata iz PDF-a i DOCX-a, sa grananjem tekst / tabela / slika.

Jedini modul u `extract_and_prepare` koji zavisi od PyMuPDF-a i python-docx-a,
i zato jedini koji se ne moze testirati bez Docker-a. Sva pravila koja se DAJU
testirati lokalno (chunkovanje, serijalizacija tabela, sklapanje manifesta)
namerno zive u `chunking.py` i `elements.py`.

Slike se ne upisuju odavde. Umesto toga se prosledjuje `image_sink` — funkcija
koja primi bajtove i vrati `image_ref`. Tako ekstrakcija ne zna za S3, pa se
moze pokrenuti i nad lokalnim fajlom sa laznim sink-om. `image_ref` je polje na
kome stoji varijanta B ablacije, pa je vazno da nastane na jednom mestu.
"""

import os
from typing import Callable, List, Optional

from elements import IMAGE, TABLE, TEXT, Element

# (bajtovi, ekstenzija bez tacke) -> image_ref
ImageSink = Callable[[bytes, str], str]

SUPPORTED_PDF = (".pdf",)
SUPPORTED_DOCX = (".docx",)


def extract(path: str, image_sink: ImageSink) -> List[Element]:
    """Bira ekstraktor po ekstenziji. Nepoznat tip je greska, ne tiho preskakanje."""
    suffix = os.path.splitext(path)[1].lower()
    if suffix in SUPPORTED_PDF:
        return extract_pdf(path, image_sink)
    if suffix in SUPPORTED_DOCX:
        return extract_docx(path, image_sink)
    raise ValueError(
        "nepodrzan format: {} (podrzani: {})".format(
            suffix or "bez ekstenzije", ", ".join(SUPPORTED_PDF + SUPPORTED_DOCX)
        )
    )


def extract_pdf(path: str, image_sink: ImageSink) -> List[Element]:
    """
    PDF preko PyMuPDF-a, stranicu po stranicu.

    Redosled je tabele -> tekst -> slike. Tabele se traze PRVE jer `get_text`
    vraca i sadrzaj celija kao obican tekst; kad bi tekst isao prvi, iste
    vrednosti bi zavrsile i kao tekstualni chunk i kao tabela, pa bi se u
    pretrazi takmicile same sa sobom.
    """
    import fitz  # PyMuPDF; uvoz je lokalan da modul mogao da se ucita bez njega

    elements: List[Element] = []
    document = fitz.open(path)
    try:
        for page_number, page in enumerate(document, start=1):
            table_bboxes = []

            tables = page.find_tables()
            for table in tables.tables:
                rows = table.extract()
                if rows:
                    table_bboxes.append(fitz.Rect(table.bbox))
                    elements.append(Element(kind=TABLE, rows=rows, page=page_number))

            text = _page_text_without(page, table_bboxes)
            if text.strip():
                elements.append(Element(kind=TEXT, text=text, page=page_number))

            for image_info in page.get_images(full=True):
                xref = image_info[0]
                extracted = document.extract_image(xref)
                ref = image_sink(extracted["image"], extracted.get("ext", "png"))
                elements.append(Element(kind=IMAGE, image_ref=ref, page=page_number))
    finally:
        document.close()

    return elements


def _page_text_without(page, exclude_rects) -> str:
    """
    Tekst stranice bez blokova koji upadaju u vec prepoznate tabele.

    Bez ovoga bi svaka tabela bila indeksirana dvaput — jednom kao Markdown,
    jednom kao razbacan tekst bez zaglavlja.
    """
    import fitz

    if not exclude_rects:
        return page.get_text("text")

    kept = []
    for block in page.get_text("blocks"):
        rect = fitz.Rect(block[:4])
        if any(rect.intersects(r) for r in exclude_rects):
            continue
        kept.append(block[4])
    return "\n".join(kept)


def extract_docx(path: str, image_sink: ImageSink) -> List[Element]:
    """
    DOCX preko python-docx-a.

    Pasusi i tabele se citaju redom kojim stoje u telu dokumenta, a ne odvojeno
    preko `doc.paragraphs` i `doc.tables`. Te dve liste gube medjusobni
    redosled, cime bi tabela ispala iz konteksta pasusa koji je uvodi.

    Uzastopni pasusi se spajaju u jedan tekstualni element, da chunkovanje ne
    bi seklo po granicama pasusa nego po 512 znakova, kako specifikacija trazi.
    """
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(path)
    elements: List[Element] = []
    buffer: List[str] = []

    def flush_text():
        if buffer:
            joined = "\n".join(buffer).strip()
            if joined:
                elements.append(Element(kind=TEXT, text=joined))
            buffer.clear()

    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            buffer.append(Paragraph(child, document).text)
        elif tag == "tbl":
            flush_text()
            table = Table(child, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            if rows:
                elements.append(Element(kind=TABLE, rows=rows))
    flush_text()

    # Slike u DOCX-u nisu vezane za poziciju u telu na nacin koji python-docx
    # lako daje, pa se citaju iz relacija dokumenta i dodaju na kraj. Posledica:
    # `page` ostaje None za slike iz DOCX-a. Prihvatljivo, jer DOCX ionako nema
    # fiksne stranice dok se ne renderuje.
    for rel in document.part.rels.values():
        if "image" not in rel.reltype:
            continue
        blob = rel.target_part.blob
        suffix = os.path.splitext(rel.target_part.partname)[1].lstrip(".") or "png"
        elements.append(Element(kind=IMAGE, image_ref=image_sink(blob, suffix)))

    return elements
